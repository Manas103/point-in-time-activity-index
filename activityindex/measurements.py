"""The five claim measurements, each an honest measured number, not a target.

Every function here returns plain Python values so scripts/run_measurements.py
and the Excel report builder can both use them without re-deriving the logic.
"""

from __future__ import annotations

import datetime
import math
import random

import pandas as pd

from activityindex import config, index_builder, oracle, publish_gate, simulate, vintage_store


def measure_structure(con) -> dict:
    """Claim 1: 85 series across the 4 named categories, genuinely combined."""
    df = con.execute(
        "SELECT category, COUNT(DISTINCT series_id) AS n FROM source_vintages GROUP BY category"
    ).fetchdf()
    counts = dict(zip(df["category"], df["n"]))
    total = int(sum(counts.values()))
    categories_present = set(counts.keys()) == set(config.CATEGORIES)
    counts_ok = all(counts.get(c, 0) == config.SERIES_PER_CATEGORY[c] for c in config.CATEGORIES)

    # "Genuine weighted combination, not a placeholder": confirm at least one
    # real release actually drew on more than one series and more than one
    # category (a placeholder would be a single hardcoded number).
    sample_release = con.execute(
        "SELECT release_date, reference_period FROM published_index WHERE lag_months = 0 LIMIT 1"
    ).fetchone()
    genuine = False
    series_used = 0
    if sample_release:
        release_date, period = sample_release
        snapshot = vintage_store.query_point_in_time(con, release_date)
        _, series_used, contributions = index_builder.compute_release_value(snapshot, period)
        categories_contributing = sum(1 for v in contributions.values() if abs(v) > 0)
        genuine = series_used > 1 and categories_contributing > 1

    value = total == 85 and categories_present and counts_ok and genuine
    return {
        "value": bool(value),
        "series_counts": counts,
        "total": total,
        "series_used_in_sample_release": series_used,
        "genuine_combination": genuine,
    }


def measure_vintage_reproduction(con, n: int = 50, oracle_n: int = 8, seed: int = config.SEED) -> dict:
    """Claim 2: recompute N published values point-in-time and diff against
    the stored value; also diff a handful against the pure-Python oracle."""
    published = con.execute(
        "SELECT release_date, reference_period, index_value, lag_months FROM published_index"
    ).fetchdf()
    rng = random.Random(seed + 2)
    n = min(n, len(published))
    sample_idx = rng.sample(range(len(published)), n)
    sample = published.iloc[sample_idx]

    matched = 0
    max_abs_diff = 0.0
    for _, row in sample.iterrows():
        as_of = config.as_date(row["release_date"])
        target_period = config.as_date(row["reference_period"])
        snapshot = vintage_store.query_point_in_time(con, as_of)
        recomputed, _count, _contrib = index_builder.compute_release_value(snapshot, target_period)
        diff = abs((recomputed if recomputed is not None else float("nan")) - row["index_value"])
        max_abs_diff = max(max_abs_diff, diff if not math.isnan(diff) else max_abs_diff)
        if recomputed is not None and math.isclose(recomputed, row["index_value"], abs_tol=1e-9):
            matched += 1

    oracle_n = min(oracle_n, len(sample))
    raw_rows = oracle.fetch_raw_rows(con)
    oracle_matched = 0
    oracle_max_abs_diff = 0.0
    for _, row in sample.iloc[:oracle_n].iterrows():
        as_of = config.as_date(row["release_date"])
        target_period = config.as_date(row["reference_period"])
        snapshot = vintage_store.query_point_in_time(con, as_of)
        fast_value, _count, _contrib = index_builder.compute_release_value(snapshot, target_period)
        oracle_value, _oracle_count = oracle.compute_index_value_oracle(raw_rows, as_of, target_period)
        diff = abs((fast_value or float("nan")) - (oracle_value if oracle_value is not None else float("nan")))
        oracle_max_abs_diff = max(oracle_max_abs_diff, diff if not math.isnan(diff) else oracle_max_abs_diff)
        if fast_value is not None and oracle_value is not None and math.isclose(fast_value, oracle_value, abs_tol=1e-9):
            oracle_matched += 1

    return {
        "matched": matched,
        "total": n,
        "max_abs_diff_stored_vs_recomputed": max_abs_diff,
        "oracle_matched": oracle_matched,
        "oracle_total": oracle_n,
        "max_abs_diff_fast_vs_oracle": oracle_max_abs_diff,
    }


def measure_revision_profile(con) -> dict:
    """Claim 3: mean/median absolute revision in index points at lag 1, 2, 3."""
    df = con.execute(
        "SELECT release_date, reference_period, lag_months, index_value FROM published_index"
    ).fetchdf()
    pivot = df.pivot_table(index="reference_period", columns="lag_months", values="index_value", aggfunc="first")
    result = {}
    for lag in (1, 2, 3):
        if 0 not in pivot.columns or lag not in pivot.columns:
            continue
        both = pivot[[0, lag]].dropna()
        abs_rev = (both[lag] - both[0]).abs()
        result[f"mean_abs_revision_{lag}mo"] = float(abs_rev.mean()) if len(abs_rev) else None
        result[f"median_abs_revision_{lag}mo"] = float(abs_rev.median()) if len(abs_rev) else None
        result[f"n_periods_{lag}mo"] = int(len(abs_rev))
    return result


def measure_publish_gate(con, all_series_ids: list[str], k_stale: int = 12, seed: int = config.SEED) -> dict:
    """Claim 4: synthetic stale scenarios must be refused and named correctly;
    real, clean historical releases must never be falsely refused."""
    rng = random.Random(seed + 4)

    # Synthetic stale scenarios: a healthy world (everyone up to date through
    # current_period) with K deliberately chosen series rolled back past the
    # staleness threshold, one scenario at a time.
    current_period = datetime.date(2024, 6, 1)
    healthy = {sid: current_period for sid in all_series_ids}
    stale_caught = 0
    for i in range(k_stale):
        rng_local = random.Random(seed + 100 + i)
        n_missing = rng_local.randint(1, 5)
        missing = rng_local.sample(all_series_ids, n_missing)
        scenario = dict(healthy)
        for sid in missing:
            # Rolled back 3 periods: beyond the 2-period staleness lookback.
            scenario[sid] = config.month_add(current_period, -3)
        result = publish_gate.evaluate(scenario, all_series_ids, current_period)
        if not result.ok and set(result.stale_series) == set(sorted(missing)):
            stale_caught += 1

    # Clean scenarios: every real historical release this project ever made.
    releases = con.execute(
        "SELECT DISTINCT release_date, reference_period FROM published_index WHERE lag_months = 0"
    ).fetchdf()
    false_refusals = 0
    m_clean = len(releases)
    for _, row in releases.iterrows():
        result = publish_gate.check(con, row["release_date"], row["reference_period"], all_series_ids)
        if not result.ok:
            false_refusals += 1

    return {
        "stale_caught": stale_caught,
        "k_stale": k_stale,
        "false_refusals": false_refusals,
        "m_clean": m_clean,
    }


def measure_excel_report(path: str) -> dict:
    """Claim 5: structural check, reopened independently with openpyxl."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path)
    except Exception as exc:  # noqa: BLE001
        return {"value": False, "error": str(exc)}
    sheet_count = len(wb.sheetnames)
    ws = wb[wb.sheetnames[0]]
    chart_count = len(getattr(ws, "_charts", []))
    ok = sheet_count == 1 and chart_count >= 1
    return {
        "value": bool(ok),
        "sheet_count": sheet_count,
        "chart_count": chart_count,
        "opened_ok": True,
    }
